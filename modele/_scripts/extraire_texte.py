"""Extraction de texte en masse : script OPTIONNEL.

L'assistant lit normalement les documents lui-même (y compris PDF et scans).
Ce script sert d'accélérateur, ou de solution de repli pour un assistant qui ne
lit pas les PDF nativement. Il est le SEUL du kit à dépendre de bibliothèques
extérieures, dont les versions sont figées dans requirements.txt, à côté de ce
fichier :

  pip install -r requirements.txt

Usage : python3 extraire_texte.py [--limite N]
Traite les fichiers en statut a_lire dont l'extension est prise en charge et
range le texte dans fichiers_contenu (limité à 50 000 caractères).
Code retour : 1 si aucune extraction n'a abouti alors que des fichiers ont échoué.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

import bdd
import scan
import sortie

LIMITE_TEXTE = 50_000


def extraire(chemin: Path) -> tuple[str, str] | None:
    """Renvoie (texte, methode), ou None si extension non prise en charge."""
    suffixe = chemin.suffix.lower()
    if suffixe == ".pdf":
        import pdfplumber

        with pdfplumber.open(chemin) as pdf:
            texte = "\n".join(page.extract_text() or "" for page in pdf.pages)
        return texte, "pdfplumber"
    if suffixe == ".docx":
        import docx

        document = docx.Document(str(chemin))
        return "\n".join(p.text for p in document.paragraphs), "python-docx"
    if suffixe == ".xlsx":
        import openpyxl

        classeur = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
        lignes = []
        for feuille in classeur.worksheets:
            for rangee in feuille.iter_rows(values_only=True):
                cellules = [str(c) for c in rangee if c is not None]
                if cellules:
                    lignes.append("\t".join(cellules))
        return "\n".join(lignes), "openpyxl"
    if suffixe in {".txt", ".md", ".csv"}:
        return chemin.read_text(encoding="utf-8", errors="replace"), "texte_brut"
    return None


def entier_positif(valeur: str) -> int:
    """0 est accepté et signifie « aucun fichier », jamais « tous »."""
    try:
        nombre = int(valeur)
    except ValueError:
        raise argparse.ArgumentTypeError(f"nombre entier attendu, reçu « {valeur} »")
    if nombre < 0:
        raise argparse.ArgumentTypeError("la limite ne peut pas être négative")
    return nombre


def main() -> None:
    sortie.configurer()
    analyseur = argparse.ArgumentParser(
        description="Extraction de texte en masse (script optionnel).",
        epilog="La qualification des documents reste le travail de l'assistant.",
    )
    analyseur.add_argument("--limite", type=entier_positif, default=None,
                           help="nombre maximum de fichiers à traiter (0 = aucun)")
    limite = analyseur.parse_args().limite

    conn = bdd.connexion()
    lignes = conn.execute(
        "SELECT f.id, f.chemin FROM fichiers f"
        " LEFT JOIN fichiers_contenu c ON c.fichier_id = f.id"
        " WHERE f.statut_classement = 'a_lire' AND f.supprime = 0 AND c.fichier_id IS NULL"
    ).fetchall()
    if limite is not None:
        lignes = lignes[:limite]

    faits = ignores = erreurs = 0
    for ligne in lignes:
        chemin = Path(ligne["chemin"])
        if scan.est_fichier_du_kit(chemin):
            ignores += 1        # mode d'emploi du kit, pas un document de l'utilisateur
            continue
        try:
            resultat = extraire(chemin)
        except Exception as erreur:
            erreurs += 1
            print(f"ERREUR  {chemin} : {erreur}")
            continue
        if resultat is None:
            ignores += 1
            continue
        texte, methode = resultat
        with conn:
            conn.execute(
                "INSERT OR REPLACE INTO fichiers_contenu (fichier_id, texte, methode, date_extraction)"
                " VALUES (?, ?, ?, ?)",
                (ligne["id"], texte[:LIMITE_TEXTE], methode, bdd.maintenant()),
            )
        faits += 1

    print(f"extraits : {faits} | non pris en charge : {ignores} | erreurs : {erreurs}")
    print("Rappel : l'extraction ne remplace pas la lecture par l'assistant :"
          " la qualification (type, émetteur, résumé…) reste son travail.")
    conn.close()
    if erreurs and not faits:
        sys.exit(1)


if __name__ == "__main__":
    main()
